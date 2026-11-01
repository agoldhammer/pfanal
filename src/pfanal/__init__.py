"""List positions from the newest portfolio export, sorted by maturity date."""

import csv
import http.cookiejar
import re
import subprocess
import sys
import urllib.parse
import urllib.request
from collections.abc import Callable, Iterator
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import NamedTuple

from openpyxl import load_workbook

# The repo's data/ directory, wherever pfanal is run from.
DATA_DIR = Path(__file__).resolve().parents[2] / "data"
# Fidelity's export name. Downloads holds unrelated files, so only these count there.
EXPORT_GLOB = "Portfolio_Positions*"

# A mm/dd/yyyy date at the very end of the description (column D).
MATURITY_RE = re.compile(r"(\d{2}/\d{2}/\d{4})\s*$")
# Date embedded in export filenames, e.g. Portfolio_Positions_Oct-17-2026.xlsx
FILENAME_DATE_RE = re.compile(r"([A-Z][a-z]{2}-\d{2}-\d{4})")

# TreasuryDirect's daily prices for all marketable Treasury securities, by CUSIP.
PRICE_URL = "https://www.treasurydirect.gov/GA-FI/FedInvest/selectSecurityPriceDate"
# Prices are published once per business day, so look back far enough to
# cover a weekend plus a holiday.
PRICE_LOOKBACK_DAYS = 7
PRICE_ROW_RE = re.compile(r"<tr>(.*?)</tr>", re.S)
PRICE_CELL_RE = re.compile(r"<td[^>]*>(.*?)</td>", re.S)


class Position(NamedTuple):
    maturity: date
    account_name: str
    symbol: str
    cost_basis: float | None  # None when Fidelity shows "--"
    present_value: float
    maturity_value: float
    priced_online: bool = False

    @property
    def gain(self) -> float | None:
        if self.cost_basis is None:
            return None
        return self.present_value - self.cost_basis

    @property
    def gain_pct(self) -> float | None:
        """Gain to date as a percentage of cost basis."""
        if self.cost_basis is None:
            return None
        return self.gain / self.cost_basis * 100


def file_date(path: Path) -> datetime:
    """Date a spreadsheet was exported: from its filename if present, else its mtime."""
    if m := FILENAME_DATE_RE.search(path.stem):
        try:
            return datetime.strptime(m.group(1), "%b-%d-%Y")
        except ValueError:
            pass
    return datetime.fromtimestamp(path.stat().st_mtime)


def windows_downloads() -> Path | None:
    """The Windows user's Downloads directory as a WSL path, or None if not on WSL."""
    try:
        # Run from a Windows drive so cmd.exe doesn't warn about a UNC working directory.
        profile = subprocess.run(
            ["cmd.exe", "/c", "echo %USERPROFILE%"],
            capture_output=True, text=True, timeout=10, cwd="/mnt/c",
        ).stdout.strip()
        home = subprocess.run(
            ["wslpath", "-u", profile], capture_output=True, text=True, timeout=10
        ).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return None
    downloads = Path(home) / "Downloads"
    return downloads if home and downloads.is_dir() else None


def spreadsheets(directory: Path, pattern: str = "*") -> list[Path]:
    """The .xlsx and .csv files in a directory, skipping Excel lock files."""
    return [
        p for ext in ("xlsx", "csv") for p in directory.glob(f"{pattern}.{ext}")
        if not p.name.startswith("~$")
    ]


def newest(files: list[Path]) -> Path:
    return max(files, key=lambda p: (file_date(p), p.stat().st_mtime))


def latest_spreadsheet(data_dir: Path = DATA_DIR) -> Path:
    if not (files := spreadsheets(data_dir)):
        raise FileNotFoundError(f"no .xlsx or .csv files found in {data_dir}/")
    return newest(files)


def find_spreadsheet() -> Path:
    """The newest Fidelity export in Windows Downloads, else the newest file in data/."""
    if (downloads := windows_downloads()) and (files := spreadsheets(downloads, EXPORT_GLOB)):
        return newest(files)
    return latest_spreadsheet()


def parse_maturity(description: object) -> date | None:
    """Return the trailing mm/dd/yyyy date in a description, or None if absent."""
    if not isinstance(description, str):
        return None
    if m := MATURITY_RE.search(description):
        return datetime.strptime(m.group(1), "%m/%d/%Y").date()
    return None


def parse_number(value: object) -> float | None:
    """A cell as a number; CSV cells look like "$99594.00" or "+$0.011", "--" if absent."""
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace("$", "").replace(",", "")
        try:
            return float(text)
        except ValueError:
            return None
    return None


def data_rows(path: Path) -> Iterator[tuple]:
    """Rows after the header, from either an .xlsx or a .csv export."""
    if path.suffix.lower() == ".csv":
        # Fidelity's CSV starts with a byte-order mark.
        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            next(reader, None)
            yield from (tuple(row) for row in reader)
    else:
        ws = load_workbook(path, read_only=True, data_only=True).active
        yield from ws.iter_rows(min_row=2, max_col=14, values_only=True)


def read_positions(path: Path) -> list[Position]:
    positions = []
    # Columns: A=account number, B=account name, C=symbol, D=description,
    # E=quantity, F=last price, G=last price change, H=current value,
    # I-M=gain/loss and percent of account, N=cost basis total.
    for row in data_rows(path):
        # Footer rows are shorter than a position row.
        row = (*row, *[None] * (14 - len(row)))
        _, account_name, symbol, description, quantity, _, _, current_value = row[:8]
        cost_basis = row[13]
        # Skips blank rows and the disclaimer footer, which have no maturity date.
        if (maturity := parse_maturity(description)) is None:
            continue
        # Bills are zero-coupon, so face value (quantity) is what's paid at maturity.
        positions.append(
            Position(
                maturity,
                str(account_name),
                str(symbol),
                parse_number(cost_basis),
                parse_number(current_value),
                parse_number(quantity),
            )
        )
    return sorted(positions, key=lambda p: p.maturity)


def parse_price_table(html: str) -> tuple[dict[str, float], str]:
    """Map CUSIP to price (per 100 face) from a TreasuryDirect price page.

    Uses the end of day price, or the sell price where end of day is still 0,
    and returns which of the two ("end of day" or "sell") was used.
    """
    prices = {}
    basis = "end of day"
    # Cells: CUSIP, security type, rate, maturity, call date, buy, sell, end of day.
    for row in PRICE_ROW_RE.findall(html):
        cells = [c.strip() for c in PRICE_CELL_RE.findall(row)]
        if len(cells) != 8:
            continue
        # End of day stays 0 until the next business day, but sell is quoted.
        if (price := float(cells[7])) > 0:
            prices[cells[0]] = price
        elif (price := float(cells[6])) > 0:
            prices[cells[0]] = price
            basis = "sell"
    return prices, basis


def treasurydirect_page(day: date) -> str:
    """Fetch TreasuryDirect's price page for a day (its table is empty if unpublished)."""
    opener = urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar())
    )
    opener.addheaders = [("User-Agent", "Mozilla/5.0")]
    # The form needs the CSRF token and session cookie from a prior GET.
    form = opener.open(PRICE_URL, timeout=30).read().decode()
    token = re.search(r'name="_csrf" value="([^"]*)"', form).group(1)
    data = urllib.parse.urlencode(
        {"priceDate": day.isoformat(), "_csrf": token, "submit": "Show Prices"}
    ).encode()
    return opener.open(PRICE_URL, data, timeout=30).read().decode()


def fetch_prices(
    today: date, get_page: Callable[[date], str] = treasurydirect_page
) -> tuple[date, dict[str, float], str]:
    """Prices from the most recent day on or before today that has them published."""
    for back in range(PRICE_LOOKBACK_DAYS):
        day = today - timedelta(days=back)
        prices, basis = parse_price_table(get_page(day))
        if prices:
            return day, prices, basis
    raise LookupError(f"no Treasury prices published in the {PRICE_LOOKBACK_DAYS} days to {today}")


def reprice(positions: list[Position], prices: dict[str, float]) -> list[Position]:
    """Value each position at its quoted price; keep the spreadsheet value if unquoted."""
    return [
        p._replace(present_value=p.maturity_value * prices[p.symbol] / 100, priced_online=True)
        if p.symbol in prices
        else p
        for p in positions
    ]


def days_until(maturity: date, today: date | None = None) -> int:
    """Calendar days from today until maturity (negative if already matured)."""
    return (maturity - (today or date.today())).days


def fmt(value: float | None, width: int, spec: str, suffix: str = "") -> str:
    """Format a number right-aligned in width, or "--" if it is missing."""
    text = "--" if value is None else f"{value:{spec}}{suffix}"
    return f"{text:>{width}}"


def main() -> None:
    try:
        path = find_spreadsheet()
    except FileNotFoundError as e:
        sys.exit(f"error: {e}")
    print(f"Source: {path}\n")
    positions = read_positions(path)
    if not positions:
        sys.exit(f"error: no positions with maturity dates found in {path}")
    today = date.today()
    try:
        price_date, prices, basis = fetch_prices(today)
    except (OSError, LookupError, AttributeError) as e:
        print(f"warning: could not fetch Treasury prices ({e}); "
              "using spreadsheet values\n", file=sys.stderr)
    else:
        print(f"Prices: TreasuryDirect {basis}, {price_date:%m/%d/%Y}\n")
        positions = reprice(positions, prices)
    width = max((len(p.account_name) for p in positions), default=12)
    print(
        f"{'Maturity':<10}  {'Days':>4}  {'Account name':<{width}}  {'Symbol':<9}  "
        f"{'Cost basis':>15}  {'Present value':>15}  {'Gain to date':>12}  "
        f"{'Value at maturity':>17}"
    )
    for p in positions:
        print(
            f"{p.maturity:%m/%d/%Y}  {days_until(p.maturity, today):>4}  "
            f"{p.account_name:<{width}}  {p.symbol:<9}  "
            f"{fmt(p.cost_basis, 15, ',.2f')}  {p.present_value:>15,.2f}  "
            f"{fmt(p.gain_pct, 12, '.3f', '%')}  "
            f"{p.maturity_value:>17,.2f}{'' if p.priced_online else '  *'}"
        )
    # Gain is measured only over positions whose cost basis is known.
    with_cb = [p for p in positions if p.cost_basis is not None]
    total_cb = sum(p.cost_basis for p in with_cb)
    total_gain_pct = (
        (sum(p.present_value for p in with_cb) - total_cb) / total_cb * 100 if total_cb else None
    )
    total_pv = sum(p.present_value for p in positions)
    total_mv = sum(p.maturity_value for p in positions)
    label_width = 10 + 2 + 4 + 2 + width + 2 + 9
    print("-" * (label_width + 2 + 15 + 2 + 15 + 2 + 12 + 2 + 17))
    print(
        f"{'Total':<{label_width}}  {total_cb:>15,.2f}  {total_pv:>15,.2f}  "
        f"{fmt(total_gain_pct, 12, '.3f', '%')}  {total_mv:>17,.2f}"
    )
    if len(with_cb) < len(positions):
        print("\n-- no cost basis reported; total gain excludes these positions")
    if not all(p.priced_online for p in positions):
        print("\n* present value from the spreadsheet (no online price found)")
